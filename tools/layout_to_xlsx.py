"""buildExportLayout_ 결과(JSON)를 xlsx 로 써 보는 검증 도구. 실제 앱은 Export.gs 가 구글시트로 만든다."""
import json, sys, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

lay = json.load(open(sys.argv[1]))
wb = openpyxl.Workbook(); ws = wb.active; ws.title = sys.argv[3] if len(sys.argv) > 3 else 'Sheet1'
thin = Side(style='thin', color='000000'); border = Border(left=thin, right=thin, top=thin, bottom=thin)
boxes = [lay[k] for k in ('tableRange', 'countRange', 'summaryRange', 'timeRange')]
def in_box(r, c): return any(b[0] <= r < b[0] + b[2] and b[1] <= c < b[1] + b[3] for b in boxes)
for r in range(lay['nrows']):
    for c in range(lay['ncols']):
        v = lay['cells'][r][c]
        if v == '' and not in_box(r + 1, c + 1): continue
        cell = ws.cell(r + 1, c + 1)
        if isinstance(v, str) and v.startswith("'"): v = v[1:]
        cell.value = v if v != '' else None
        cell.font = Font(name='맑은 고딕', size=35 if (r, c) == (0, 0) else 10, bold=lay['bold'][r][c], color=(lay['fg'][r][c] or '#000000')[1:].upper())
        if lay['bg'][r][c]: cell.fill = PatternFill('solid', fgColor=lay['bg'][r][c][1:].upper())
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        if in_box(r + 1, c + 1): cell.border = border
for m in lay['merges']:
    ws.merge_cells(start_row=m[0], start_column=m[1], end_row=m[0] + m[2] - 1, end_column=m[1] + m[3] - 1)
for c, w in lay['colWidths'].items(): ws.column_dimensions[get_column_letter(int(c))].width = w
for r, h in lay['rowHeights'].items(): ws.row_dimensions[int(r)].height = h
ws.freeze_panes = ws.cell(lay['freeze'][0] + 1, lay['freeze'][1] + 1)
wb.save(sys.argv[2])
